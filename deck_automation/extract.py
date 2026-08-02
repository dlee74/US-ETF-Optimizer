"""Read current portfolio holdings out of an existing Wiley Group deck.

Two sources of "current weight" exist per portfolio block and must agree:
  1. The chart's cached category/value data (ppt/charts/chartN.xml) — this is
     what PowerPoint actually renders, and the primary source of truth.
  2. The holdings table on the same slide (Holding/Weight columns).

A third piece of text exists in a summary box — a "Target Weights: ..." line
— but that is the client's strategic/IPS target allocation, a DIFFERENT
concept from current actual weight (which drifts from target as markets
move). It is parsed into its own dict (target_weights) for use by the
"refresh" pipeline mode (the metrics table is computed AT this target, per
its own subtitle text), but it is NEVER cross-validated against (1)/(2) —
current and target are legitimately different numbers — and the patcher
never writes back to this text box.

Extraction hard-fails (no silent fallback) if the chart and table disagree
beyond a small rounding tolerance, or if a ticker isn't in the client's
ticker_map.
"""

from __future__ import annotations

import re
from dataclasses import dataclass

from pptx import Presentation

WEIGHT_TOLERANCE = 0.001  # 0.1 percentage point


@dataclass
class PortfolioHoldings:
    name: str
    slide_title: str
    table_slide: int          # 1-indexed
    chart_part: str
    tickers: list[str]        # bare tickers, deck order
    weights: list[float]      # fractional, same order as tickers (current actual)
    target_weights_text: str | None       # raw "Target Weights: ..." line, for display
    target_weights: dict[str, float] | None  # parsed {ticker: fraction}, for refresh mode


def _find_chart_shape(slide, chart_part_name: str):
    for shape in slide.shapes:
        if not shape.has_chart:
            continue
        partname = str(shape.chart.part.partname)  # e.g. "/ppt/charts/chart2.xml"
        if partname.endswith("/" + chart_part_name):
            return shape.chart
    raise ValueError(f"No chart shape found matching part {chart_part_name!r} "
                      f"on slide {slide.slide_id}")


def _extract_chart_holdings(slide, chart_part_name: str) -> dict[str, float]:
    chart = _find_chart_shape(slide, chart_part_name)
    plot = chart.plots[0]
    categories = list(plot.categories)
    series = plot.series[0]
    values = list(series.values)
    if len(categories) != len(values):
        raise ValueError(f"Chart {chart_part_name}: {len(categories)} categories "
                          f"!= {len(values)} values")
    return dict(zip(categories, values))


def _find_table_shape(slide):
    for shape in slide.shapes:
        if shape.has_table:
            return shape.table
    raise ValueError(f"No table found on slide {slide.slide_id}")


def _extract_table_holdings(slide) -> dict[str, float]:
    table = _find_table_shape(slide)
    header = [c.text.strip() for c in table.rows[0].cells]
    try:
        ticker_col = header.index("Holding")
        weight_col = header.index("Weight")
    except ValueError as e:
        raise ValueError(f"Expected 'Holding'/'Weight' columns, got {header}") from e

    holdings = {}
    for row in list(table.rows)[1:]:
        ticker = row.cells[ticker_col].text.strip()
        weight_text = row.cells[weight_col].text.strip()
        if ticker == "Portfolio Total" or not ticker:
            continue
        # "31.7%" -> 0.317
        m = re.match(r"([\d.]+)%", weight_text)
        if not m:
            raise ValueError(f"Could not parse weight {weight_text!r} for {ticker!r}")
        holdings[ticker] = float(m.group(1)) / 100.0
    return holdings


def _extract_target_weights_text(slide) -> str | None:
    for shape in slide.shapes:
        if not shape.has_text_frame:
            continue
        for para in shape.text_frame.paragraphs:
            text = "".join(run.text for run in para.runs)
            if text.strip().startswith("Target Weights:"):
                return text.strip()
    return None


def _parse_target_weights(text: str | None) -> dict[str, float] | None:
    """"Target Weights: HDIV 30% / ZCS 30% / RPF 15% / XFR 25%" -> fractions."""
    if text is None:
        return None
    body = text.split(":", 1)[1]
    weights = {}
    for part in body.split("/"):
        m = re.match(r"\s*(\w+)\s+([\d.]+)%\s*$", part)
        if not m:
            raise ValueError(f"Could not parse target-weight segment {part!r} in {text!r}")
        weights[m.group(1)] = float(m.group(2)) / 100.0
    return weights


def extract_portfolio_holdings(pptx_path: str, client_config: dict) -> list[PortfolioHoldings]:
    """One PortfolioHoldings per entry in client_config["portfolios"], with
    chart/table cross-validation and ticker_map coverage enforced."""
    prs = Presentation(pptx_path)
    ticker_map = client_config["ticker_map"]
    results = []

    for portfolio in client_config["portfolios"]:
        slide = prs.slides[portfolio["table_slide"] - 1]  # slides are 0-indexed

        chart_holdings = _extract_chart_holdings(slide, portfolio["chart_part"])
        table_holdings = _extract_table_holdings(slide)

        if set(chart_holdings) != set(table_holdings):
            raise ValueError(
                f"{portfolio['name']}: chart tickers {sorted(chart_holdings)} != "
                f"table tickers {sorted(table_holdings)}"
            )
        for ticker, chart_w in chart_holdings.items():
            table_w = table_holdings[ticker]
            if abs(chart_w - table_w) > WEIGHT_TOLERANCE:
                raise ValueError(
                    f"{portfolio['name']}/{ticker}: chart weight {chart_w} disagrees "
                    f"with table weight {table_w} by more than {WEIGHT_TOLERANCE}"
                )

        unmapped = set(chart_holdings) - set(ticker_map)
        if unmapped:
            raise ValueError(
                f"{portfolio['name']}: ticker(s) {sorted(unmapped)} not in ticker_map — "
                "refusing to guess an exchange suffix. Add them to clients/wiley_group.py "
                "if this is a genuine new holding, not a data error."
            )

        tickers = list(chart_holdings)  # preserves chart's cache order
        weights = [chart_holdings[t] for t in tickers]

        target_text = _extract_target_weights_text(slide)
        target_weights = _parse_target_weights(target_text)
        if target_weights is not None:
            unmapped_target = set(target_weights) - set(ticker_map)
            if unmapped_target:
                raise ValueError(
                    f"{portfolio['name']}: target-weight ticker(s) "
                    f"{sorted(unmapped_target)} not in ticker_map"
                )

        results.append(PortfolioHoldings(
            name=portfolio["name"],
            slide_title=portfolio["slide_title"],
            table_slide=portfolio["table_slide"],
            chart_part=portfolio["chart_part"],
            tickers=tickers,
            weights=weights,
            target_weights_text=target_text,
            target_weights=target_weights,
        ))

    return results
