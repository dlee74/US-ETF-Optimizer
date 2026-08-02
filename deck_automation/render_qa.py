"""QA gates for a patched deck, run in order — each is a real check, not a
rubber stamp, and any failure stops the pipeline before the file is
considered usable:

  1. Structural: zip integrity + XML well-formedness on every part this
     pipeline touched.
  2. Round-trip semantic: re-open the patched file with python-pptx
     (read-only) and confirm the metrics table actually shows what was
     computed — catches a patch that "succeeded" but wrote the wrong cell.
  3. Text-length heuristic: flag (not block) any cell whose new text is
     longer than what was there before — a cheap proxy for the wrapping/
     overlap bugs that bit the manual process twice.
  4. Visual render: PowerPoint COM exports the changed slides to PNG at a
     fixed resolution, using the actual installed PowerPoint (pixel-exact,
     no LibreOffice approximation).

The final gate is always a human/Claude-vision look at the PNGs from step 4
— this module does not auto-finalize a file. Structural checks alone did
not catch the real layout bugs seen before; nothing here should be trusted
to replace that look.
"""

from __future__ import annotations

import os
import zipfile
from dataclasses import dataclass, field

from lxml import etree
from pptx import Presentation


@dataclass
class QAReport:
    structural_ok: bool = False
    structural_errors: list[str] = field(default_factory=list)
    semantic_ok: bool = False
    semantic_errors: list[str] = field(default_factory=list)
    length_warnings: list[str] = field(default_factory=list)
    rendered_pngs: list[str] = field(default_factory=list)

    @property
    def passed_automated_checks(self) -> bool:
        return self.structural_ok and self.semantic_ok


def check_structural(pptx_path: str, touched_parts: list[str]) -> tuple[bool, list[str]]:
    errors = []
    zf = zipfile.ZipFile(pptx_path)
    bad = zf.testzip()
    if bad is not None:
        errors.append(f"Zip integrity failed at {bad!r}")

    for part in touched_parts:
        try:
            etree.fromstring(zf.read(part))
        except etree.XMLSyntaxError as e:
            errors.append(f"{part}: not well-formed XML: {e}")
    return (len(errors) == 0), errors


def check_semantic(pptx_path: str, client_config: dict, refreshed: list) -> tuple[bool, list[str]]:
    """Re-open with python-pptx and confirm the metrics table's Value column
    matches what was computed, cell by cell."""
    errors = []
    prs = Presentation(pptx_path)

    for portfolio, result in zip(client_config["portfolios"], refreshed):
        slide = prs.slides[portfolio["metrics_slide"] - 1]
        table = next((s.table for s in slide.shapes if s.has_table), None)
        if table is None:
            errors.append(f"{portfolio['name']}: no table found on metrics slide")
            continue
        found = {row.cells[0].text.strip(): row.cells[1].text.strip()
                 for row in list(table.rows)[1:]}
        from deck_automation.patch.table_text import LABEL_ALIASES
        for label, deck_value in found.items():
            metric_key = LABEL_ALIASES.get(label, label)
            if metric_key not in result.metrics.index:
                continue
            expected = str(result.metrics.loc[metric_key, "Value"])
            if deck_value != expected:
                errors.append(
                    f"{portfolio['name']}/{label}: deck shows {deck_value!r}, "
                    f"expected {expected!r}"
                )
    return (len(errors) == 0), errors


def check_text_length(before_pptx_path: str, after_pptx_path: str,
                       client_config: dict) -> list[str]:
    """Cheap proxy for wrap/overlap bugs: any cell whose text got LONGER
    (not just changed) is worth a second look before finalizing."""
    warnings = []
    before = Presentation(before_pptx_path)
    after = Presentation(after_pptx_path)

    for portfolio in client_config["portfolios"]:
        idx = portfolio["metrics_slide"] - 1
        before_table = next((s.table for s in before.slides[idx].shapes if s.has_table), None)
        after_table = next((s.table for s in after.slides[idx].shapes if s.has_table), None)
        if before_table is None or after_table is None:
            continue
        for b_row, a_row in zip(list(before_table.rows)[1:], list(after_table.rows)[1:]):
            b_text, a_text = b_row.cells[1].text, a_row.cells[1].text
            if len(a_text) > len(b_text):
                warnings.append(
                    f"{portfolio['name']}/{b_row.cells[0].text}: "
                    f"{b_text!r} ({len(b_text)} chars) -> {a_text!r} ({len(a_text)} chars)"
                )
    return warnings


def render_slides_to_png(pptx_path: str, slide_numbers: list[int],
                          out_dir: str, width: int = 1920, height: int = 1080) -> list[str]:
    """PowerPoint COM automation — Windows only. Requires PowerPoint
    installed (confirmed present on this machine). Renders 1-indexed
    slide_numbers, matching how client configs talk about slides."""
    import win32com.client

    os.makedirs(out_dir, exist_ok=True)
    pptx_path = os.path.abspath(pptx_path)
    out_dir = os.path.abspath(out_dir)

    app = win32com.client.Dispatch("PowerPoint.Application")
    pngs = []
    pres = app.Presentations.Open(pptx_path, WithWindow=False)
    try:
        for n in slide_numbers:
            out_png = os.path.join(out_dir, f"slide{n}.png")
            pres.Slides(n).Export(out_png, "PNG", width, height)
            pngs.append(out_png)
    finally:
        pres.Close()
        app.Quit()
    return pngs


def run_qa(before_pptx_path: str, after_pptx_path: str, client_config: dict,
           refreshed: list, png_out_dir: str) -> QAReport:
    report = QAReport()

    touched_parts = []
    for portfolio in client_config["portfolios"]:
        prs = Presentation(after_pptx_path)
        touched_parts.append(str(prs.slides[portfolio["metrics_slide"] - 1].part.partname).lstrip("/"))

    report.structural_ok, report.structural_errors = check_structural(after_pptx_path, touched_parts)
    report.semantic_ok, report.semantic_errors = check_semantic(after_pptx_path, client_config, refreshed)
    report.length_warnings = check_text_length(before_pptx_path, after_pptx_path, client_config)

    if report.passed_automated_checks:
        slide_numbers = [p["metrics_slide"] for p in client_config["portfolios"]]
        report.rendered_pngs = render_slides_to_png(after_pptx_path, slide_numbers, png_out_dir)

    return report
