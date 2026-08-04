"""python-docx writer for the trade-blotter one-pager.

Generated fresh each time (no existing branded template to preserve, unlike
the .pptx pipeline — this is a plain functional document), so there's no
chart-cache-sync-style risk here. Boilerplate text (Execution Protocol,
Disclaimer, the quick-reference callout) is copied verbatim from the real
reference document (Clients/DJ/2026/July 24/Wiley_Trade_Blotter_one_page.docx)
— it's fixed wording, not something this pipeline generates.
"""

from __future__ import annotations

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

QUICK_REFERENCE_CALLOUT = (
    "Use Day limit orders only. Enter near the bid/ask midpoint. Cap adjustments at "
    "$0.02 total. Pause and let orders expire if the spread exceeds roughly 2x its "
    "recent normal range."
)

BEFORE_EACH_TRADE_NOTE = (
    "*BEFORE EACH TRADE: Verify current prices, share quantities and spreads using "
    "live market quotes. Limit prices and ranges reflect Yahoo Finance daily closes "
    "as at {as_of_date}."
)

EXECUTION_PROTOCOL = [
    ("ORDER", "Use Day limit orders only. Avoid market orders unless specifically "
               "authorized; they may execute at an unfavorable price. Confirm expiry "
               "at today's close (not GTC)."),
    ("TIMING", "Preferred trading window: 10:30 a.m. to 2:00 p.m. ET. When practical, "
                "avoid the first and final 15 minutes of the trading day."),
    ("PRICE", "Enter near the midpoint of the bid-ask spread. Do not adjust the limit "
               "more than $0.02 in total; allowing the order to expire is generally "
               "preferred over chasing the price."),
    ("UNFILLED / PARTIAL", "Leave any unfilled or partially filled order active until "
                            "market close. Record partial fills and allow the remaining "
                            "balance to expire unless new instructions are provided."),
    ("HALT TRIGGERS", "If the bid-ask spread exceeds roughly 2x its recent normal "
                       "range, pause trading in that ETF. If significant news or "
                       "unusual volatility occurs, cancel open orders & consult the "
                       "advisor."),
]

AFTER_THE_SESSION_NOTE = (
    "AFTER THE SESSION: Record completed trades, partial fills, expiries and observed "
    "bid-ask spreads. Review significant overnight market developments. If an ETF "
    "remains unfilled for two consecutive trading days, notify the advisor before "
    "another attempt. Do not combine missed orders without updated instructions."
)

DISCLAIMER = (
    "DISCLAIMER  Execution reference only; not investment advice or a recommendation. "
    "Verify all prices, quantities, spreads and liquidity at trade time. This protocol "
    "reflects conditions as of {as_of_date} and must be reassessed if conditions "
    "change. InvestMint Inc. does not custody assets or execute trades. All investment "
    "and execution decisions remain the responsibility of the account owner and their "
    "registered investment advisor or broker, as applicable."
)

TRADE_TABLE_COLUMNS = ["Ticker", "Action", "Order Type", "Shares", "Amount", "Limit Price",
                        "Recent Range (~9 mo)", "New Weight"]
WATCHPOINTS_COLUMNS = ["Ticker", "Current market considerations", "Execution guidance"]


def _add_table(doc, rows: list[list[str]], header: bool = True):
    table = doc.add_table(rows=len(rows), cols=len(rows[0]))
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j, cell_text in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(cell_text)
            if header and i == 0:
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.bold = True
    return table


def build_blotter(client_name: str, as_of_date: str,
                   portfolios: list[dict], out_path: str) -> None:
    """portfolios: list of {"heading": str, "rows": list[dict] (from
    trade_blotter.build_blotter_row(), TRADE_TABLE_COLUMNS keys),
    "watchpoints": list of {"Ticker", "Current market considerations",
    "Execution guidance"}} — watchpoints text is supplied by the caller
    (Claude, reviewing the computed rows), not generated here."""
    doc = Document()

    title = doc.add_heading(f"{client_name.upper()} CORPORATE CASH PORTFOLIOS", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run(f"ONE-PAGE TRADE BLOTTER & EXECUTION INSTRUCTIONS  |  As of {as_of_date}")
    run.bold = True

    doc.add_heading("TRADE INSTRUCTIONS", level=2)

    callout = doc.add_paragraph()
    callout.add_run(QUICK_REFERENCE_CALLOUT).italic = True

    for portfolio in portfolios:
        doc.add_paragraph(portfolio["heading"]).runs[0].bold = True
        rows = [TRADE_TABLE_COLUMNS] + [
            [r[col] for col in TRADE_TABLE_COLUMNS] for r in portfolio["rows"]
        ]
        _add_table(doc, rows)

    note = doc.add_paragraph()
    note.add_run(BEFORE_EACH_TRADE_NOTE.format(as_of_date=as_of_date)).italic = True

    doc.add_heading("EXECUTION PROTOCOL", level=2)
    _add_table(doc, [[k, v] for k, v in EXECUTION_PROTOCOL], header=False)

    all_watchpoints = [row for p in portfolios for row in p.get("watchpoints", [])]
    if all_watchpoints:
        doc.add_heading("TICKER-SPECIFIC WATCHPOINTS", level=2)
        rows = [WATCHPOINTS_COLUMNS] + [
            [w[col] for col in WATCHPOINTS_COLUMNS] for w in all_watchpoints
        ]
        _add_table(doc, rows)

    doc.add_paragraph(AFTER_THE_SESSION_NOTE)

    disclaimer = doc.add_paragraph()
    disclaimer_run = disclaimer.add_run(DISCLAIMER.format(as_of_date=as_of_date))
    disclaimer_run.font.size = Pt(8)

    doc.save(out_path)
